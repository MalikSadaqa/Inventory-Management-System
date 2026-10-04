from __future__ import annotations

from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook


def test_create_invoice_computes_totals_with_16_percent_tax(client, api):
    category = api.category()
    pen = api.item(category["id"], name="Pen", price="10.00")
    pad = api.item(category["id"], name="Pad", price="2.50")
    customer = api.customer()

    invoice = api.invoice(customer["id"], [(pen["id"], 3), (pad["id"], 2)])

    assert invoice["total_quantity"] == 5
    assert Decimal(invoice["subtotal"]) == Decimal("35.00")
    assert Decimal(invoice["tax_rate"]) == Decimal("0.16")
    assert Decimal(invoice["tax_amount"]) == Decimal("5.60")
    assert Decimal(invoice["total"]) == Decimal("40.60")
    assert [line["line_subtotal"] for line in invoice["lines"]] == ["30.00", "5.00"]


def test_tax_is_rounded_half_up_to_cents(client, api):
    item = api.item(api.category()["id"], price="0.31")

    invoice = api.invoice(api.customer()["id"], [(item["id"], 1)])

    # 0.31 * 0.16 = 0.0496
    assert Decimal(invoice["tax_amount"]) == Decimal("0.05")
    assert Decimal(invoice["total"]) == Decimal("0.36")


def test_invoice_number_is_derived_from_invoice_id(client, api):
    item = api.item(api.category()["id"])
    customer = api.customer()

    first = api.invoice(customer["id"], [(item["id"], 1)])
    second = api.invoice(customer["id"], [(item["id"], 1)])

    assert first["invoice_number"] == f"INV-{first['id']:06d}"
    assert second["invoice_number"] == f"INV-{second['id']:06d}"


def test_invoice_lines_keep_snapshot_after_item_changes(client, api):
    root = api.category("Stickers")
    leaf = api.category("Round", parent_id=root["id"])
    item = api.item(leaf["id"], name="Sticker", price="5.00", cost="1.00")
    invoice = api.invoice(api.customer()["id"], [(item["id"], 2)])

    client.patch(f"/items/{item['id']}", json={"name": "Renamed", "price": "99.00"})
    line = client.get(f"/invoices/{invoice['id']}").json()["lines"][0]

    assert line["item_name_snapshot"] == "Sticker"
    assert line["unit_price_snapshot"] == "5.00"
    assert line["unit_cost_snapshot"] == "1.00"
    assert line["category_path_snapshot"] == "Stickers / Round"
    assert line["line_subtotal"] == "10.00"


def test_create_invoice_for_missing_customer_returns_404(client, api):
    item = api.item(api.category()["id"])

    response = client.post(
        "/invoices",
        json={"customer_id": 999999, "invoice_date": "2026-01-15", "items": [{"item_id": item["id"], "quantity": 1}]},
    )

    assert response.status_code == 404
    assert response.json()["detail"] == "Customer not found."


def test_create_invoice_with_missing_item_returns_404(client, api):
    customer = api.customer()

    response = client.post(
        "/invoices",
        json={"customer_id": customer["id"], "invoice_date": "2026-01-15", "items": [{"item_id": 999999, "quantity": 1}]},
    )

    assert response.status_code == 404
    assert response.json()["detail"] == "Item not found: 999999."


def test_create_invoice_requires_lines_and_positive_quantities(client, api):
    customer = api.customer()
    item = api.item(api.category()["id"])
    base = {"customer_id": customer["id"], "invoice_date": "2026-01-15"}

    no_lines = client.post("/invoices", json={**base, "items": []})
    zero_quantity = client.post("/invoices", json={**base, "items": [{"item_id": item["id"], "quantity": 0}]})

    assert no_lines.status_code == 422
    assert zero_quantity.status_code == 422


def test_list_invoices_newest_first_and_paginated(client, api):
    item = api.item(api.category()["id"])
    customer = api.customer()
    for day in ["2026-01-01", "2026-03-01", "2026-02-01"]:
        api.invoice(customer["id"], [(item["id"], 1)], invoice_date=day)

    first_page = client.get("/invoices", params={"page_size": 2}).json()
    second_page = client.get("/invoices", params={"page_size": 2, "page": 2}).json()

    assert first_page["total"] == 3
    assert [i["invoice_date"] for i in first_page["items"]] == ["2026-03-01", "2026-02-01"]
    assert [i["invoice_date"] for i in second_page["items"]] == ["2026-01-01"]


def test_get_missing_invoice_returns_404(client):
    for path in ["/invoices/999999", "/invoices/999999/pdf", "/invoices/999999/pdf/download", "/invoices/999999/excel"]:
        assert client.get(path).status_code == 404, path


def test_invoice_pdf_inline_and_download(client, api):
    item = api.item(api.category()["id"])
    invoice = api.invoice(api.customer()["id"], [(item["id"], 1)])
    filename = f'{invoice["invoice_number"]}.pdf'

    inline = client.get(f"/invoices/{invoice['id']}/pdf")
    download = client.get(f"/invoices/{invoice['id']}/pdf/download")

    assert inline.status_code == 200
    assert inline.headers["content-type"] == "application/pdf"
    assert inline.headers["content-disposition"] == f'inline; filename="{filename}"'
    assert inline.content.startswith(b"%PDF")
    assert download.headers["content-disposition"] == f'attachment; filename="{filename}"'


def test_invoice_excel_export_has_numeric_amounts(client, api):
    item = api.item(api.category()["id"], name="Pen", price="10.00")
    invoice = api.invoice(api.customer("Ann")["id"], [(item["id"], 3)])

    response = client.get(f"/invoices/{invoice['id']}/excel")

    assert response.status_code == 200
    assert response.headers["content-disposition"] == f'attachment; filename="{invoice["invoice_number"]}.xlsx"'
    workbook = load_workbook(BytesIO(response.content))
    summary = {row[0].value: row[1] for row in workbook["Invoice"].iter_rows(min_row=2)}
    assert summary["Invoice Number"].value == invoice["invoice_number"]
    assert summary["Customer Name"].value == "Ann"
    assert summary["Total"].value == 34.8
    assert summary["Total"].number_format == "#,##0.00"
    assert summary["Tax Rate"].number_format == "0.00%"

    line = [cell.value for cell in next(workbook["Line Items"].iter_rows(min_row=2))]
    assert line[0] == "Pen"
    assert line[3:] == [3, 10, 30]
